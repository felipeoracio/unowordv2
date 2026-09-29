"""Validated text-layer extraction for UnoWord document uploads."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_DOCX_FILES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024

MIME_BY_EXTENSION = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class ExtractedDocument:
    filename: str
    file_type: str
    text: str
    word_count: int


class DocumentProcessingError(ValueError):
    pass


def count_words(text: str) -> int:
    return sum(1 for token in (text or "").split() if any(character.isalnum() for character in token))


def _extract_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("TXT and Markdown files must use UTF-8 encoding") from exc


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise DocumentProcessingError("The file does not contain a valid PDF header")
    try:
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted and reader.decrypt("") == 0:
            raise DocumentProcessingError("Password-protected PDFs are not supported")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentProcessingError(f"PDFs are limited to {MAX_PDF_PAGES} pages")
        return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError("The PDF could not be read safely") from exc


def _validate_docx_archive(data: bytes) -> None:
    if not data.startswith(b"PK\x03\x04"):
        raise DocumentProcessingError("The file does not contain a valid DOCX archive")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_FILES:
                raise DocumentProcessingError("The DOCX archive contains too many files")
            if sum(member.file_size for member in members) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentProcessingError("The DOCX expands beyond the safe processing limit")
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentProcessingError("The DOCX structure is incomplete")
            for name in names:
                if name.startswith("/") or ".." in PurePosixPath(name).parts:
                    raise DocumentProcessingError("The DOCX contains an unsafe archive path")
    except BadZipFile as exc:
        raise DocumentProcessingError("The DOCX archive is invalid") from exc


def _extract_docx(data: bytes) -> str:
    _validate_docx_archive(data)
    try:
        document = Document(BytesIO(data))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                text = "\t".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    blocks.append(text)
        return "\n\n".join(blocks)
    except Exception as exc:
        raise DocumentProcessingError("The DOCX could not be read safely") from exc


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    safe_filename = Path(filename or "").name
    extension = Path(safe_filename).suffix.lower()
    if extension not in MIME_BY_EXTENSION:
        raise DocumentProcessingError("Supported formats are TXT, Markdown, PDF, and DOCX")
    if not data:
        raise DocumentProcessingError("The uploaded file is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentProcessingError("The uploaded file exceeds the 5 MB limit")

    if extension in {".txt", ".md", ".markdown"}:
        text = _extract_text(data)
    elif extension == ".pdf":
        text = _extract_pdf(data)
    else:
        text = _extract_docx(data)

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentProcessingError("No extractable text was found; scanned documents require a future OCR processor")
    return ExtractedDocument(
        filename=safe_filename,
        file_type=MIME_BY_EXTENSION[extension],
        text=text,
        word_count=count_words(text),
    )